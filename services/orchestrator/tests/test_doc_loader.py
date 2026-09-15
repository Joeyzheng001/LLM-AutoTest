"""文档加载器测试：md/txt 直读、docx/xlsx 文本化、不支持格式拒绝。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from orchestrator.doc_loader import load_document_text
from orchestrator.errors import PipelineError


def test_load_md(tmp_path: Path) -> None:
    """md/txt 按原文读取（保持行序）。"""
    doc = tmp_path / "需求.md"
    doc.write_text("# 标题\n\n功能 A：查询用户。", encoding="utf-8")
    text = load_document_text(doc)
    assert "查询用户" in text
    assert text.splitlines()[0] == "# 标题"


def test_load_docx(tmp_path: Path) -> None:
    """docx 段落 + 表格文本化（表格行以制表符拼接）。"""
    document = Document()
    document.add_paragraph("功能 A：按 ID 查询用户")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "编号"
    table.cell(0, 1).text = "需求"
    table.cell(1, 0).text = "R1"
    table.cell(1, 1).text = "导出报表"
    doc = tmp_path / "需求.docx"
    document.save(str(doc))

    text = load_document_text(doc)
    assert "按 ID 查询用户" in text
    assert "编号\t需求" in text
    assert "R1\t导出报表" in text


def test_load_xlsx(tmp_path: Path) -> None:
    """xlsx 以 sheet 名作节标题、行制表符拼接。"""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "需求清单"
    sheet.append(["编号", "需求"])
    sheet.append(["R1", "用户登录"])
    doc = tmp_path / "需求.xlsx"
    workbook.save(str(doc))

    text = load_document_text(doc)
    assert "## 需求清单" in text
    assert "R1\t用户登录" in text


def test_missing_file_rejected(tmp_path: Path) -> None:
    """文件不存在 → X-ORCH-01。"""
    with pytest.raises(PipelineError) as excinfo:
        load_document_text(tmp_path / "不存在.docx")
    assert excinfo.value.code == "X-ORCH-01"


def test_unsupported_format_rejected(tmp_path: Path) -> None:
    """不支持格式（pdf 等）→ X-ORCH-01，显式列出支持集。"""
    doc = tmp_path / "需求.pdf"
    doc.write_text("fake", encoding="utf-8")
    with pytest.raises(PipelineError, match="不支持的需求文档格式"):
        load_document_text(doc)
